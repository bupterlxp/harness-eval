"""Context management for data analysis tasks.

Handles data discovery, schema inference, compact summaries,
and prompt construction within token budgets.
"""
from __future__ import annotations

import csv
import json
import os
from pathlib import Path
from typing import Any


# File patterns to discover in the workdir
DATA_EXTENSIONS = {
    ".csv", ".tsv", ".json", ".jsonl", ".parquet", ".xlsx",
    ".xls", ".sqlite", ".db", ".sql",
}

DATA_NAME_PATTERNS = (
    "train", "test", "sample_submission", "submission",
    "data", "dataset", "features", "labels", "target",
)

INFO_NAME_PATTERNS = (
    "readme", "instructions", "description", "overview",
    "data_description", "column_description",
)


def discover_files(workdir: Path) -> dict[str, list[Path]]:
    """Discover data, info, and other files in the workdir."""
    result: dict[str, list[Path]] = {
        "data_files": [],
        "info_files": [],
        "other_files": [],
        "sample_submission": None,  # type: ignore
    }
    workdir = Path(workdir)
    if not workdir.exists():
        return result

    for p in sorted(workdir.rglob("*")):
        if not p.is_file():
            continue
        if any(skip in p.parts for skip in {
            ".git", "__pycache__", ".venv", "node_modules",
            ".ipynb_checkpoints", "output", "dev_bmk_runs",
        }):
            continue

        name_lower = p.name.lower()
        suffix = p.suffix.lower()

        if name_lower.startswith("sample_submission"):
            result["sample_submission"] = p
            result["data_files"].append(p)
        elif suffix in DATA_EXTENSIONS:
            result["data_files"].append(p)
        elif any(pat in name_lower for pat in INFO_NAME_PATTERNS):
            result["info_files"].append(p)
        elif suffix in {".py", ".ipynb", ".r", ".rmd"}:
            result["other_files"].append(p)

    return result


def infer_schema(filepath: Path, max_rows: int = 5) -> dict[str, Any]:
    """Infer schema and preview of a data file."""
    filepath = Path(filepath)
    suffix = filepath.suffix.lower()
    info: dict[str, Any] = {
        "path": str(filepath),
        "name": filepath.name,
        "extension": suffix,
        "size_bytes": filepath.stat().st_size if filepath.exists() else 0,
    }

    try:
        if suffix in {".csv", ".tsv"}:
            info.update(_csv_schema(filepath, max_rows))
        elif suffix in {".json", ".jsonl"}:
            info.update(_json_schema(filepath, max_rows))
        elif suffix == ".parquet":
            info.update(_parquet_schema(filepath))
        elif suffix in {".sqlite", ".db"}:
            info.update(_sqlite_schema(filepath))
        elif suffix in {".xlsx", ".xls"}:
            info.update(_excel_schema(filepath))
    except Exception as exc:
        info["schema_error"] = str(exc)

    return info


def _csv_schema(filepath: Path, max_rows: int) -> dict[str, Any]:
    """Extract schema and preview from CSV/TSV."""
    sep = "\t" if filepath.suffix.lower() == ".tsv" else ","
    info: dict[str, Any] = {}

    with filepath.open(newline="", encoding="utf-8", errors="replace") as f:
        reader = csv.reader(f, delimiter=sep)
        try:
            header = next(reader)
        except StopIteration:
            return {"columns": [], "row_count": 0, "preview": []}

        info["columns"] = header
        rows = []
        for i, row in enumerate(reader):
            if i >= max_rows:
                break
            rows.append(row)
        info["preview"] = rows
        info["preview_row_count"] = len(rows)

    # Count total rows
    with filepath.open(newline="", encoding="utf-8", errors="replace") as f:
        row_count = sum(1 for _ in f) - 1  # subtract header
    info["row_count"] = max(0, row_count)

    return info


def _json_schema(filepath: Path, max_rows: int) -> dict[str, Any]:
    """Extract schema from JSON/JSONL."""
    info: dict[str, Any] = {}
    with filepath.open(encoding="utf-8", errors="replace") as f:
        content = f.read(500_000)  # cap read

    if filepath.suffix.lower() == ".jsonl":
        lines = content.strip().split("\n")
        rows = []
        for i, line in enumerate(lines[:max_rows + 1]):
            if not line.strip():
                continue
            try:
                rows.append(json.loads(line))
            except json.JSONDecodeError:
                pass
        info["total_lines"] = len(lines)
        info["preview"] = rows[:max_rows]
        if rows and isinstance(rows[0], dict):
            info["columns"] = list(rows[0].keys())
            info["row_count"] = len(lines)
    else:
        try:
            data = json.loads(content)
            if isinstance(data, list):
                info["row_count"] = len(data)
                info["preview"] = data[:max_rows]
                if data and isinstance(data[0], dict):
                    info["columns"] = list(data[0].keys())
            elif isinstance(data, dict):
                info["type"] = "object"
                info["keys"] = list(data.keys())[:50]
                info["preview"] = {k: v for k, v in list(data.items())[:10]}
        except json.JSONDecodeError as exc:
            info["json_error"] = str(exc)

    return info


def _parquet_schema(filepath: Path) -> dict[str, Any]:
    """Extract schema from parquet using minimal import."""
    try:
        import pyarrow.parquet as pq
        table = pq.read_table(filepath, memory_map=True)
        return {
            "columns": table.column_names,
            "row_count": table.num_rows,
            "dtypes": {str(table.schema.field(c).type) for c in table.column_names},
        }
    except ImportError:
        return {"schema_error": "pyarrow not available for parquet"}


def _sqlite_schema(filepath: Path) -> dict[str, Any]:
    """Extract schema from SQLite database."""
    try:
        import sqlite3
        conn = sqlite3.connect(str(filepath))
        cursor = conn.cursor()
        cursor.execute("SELECT name FROM sqlite_master WHERE type='table'")
        tables = [row[0] for row in cursor.fetchall()]
        info: dict[str, Any] = {"tables": {}}
        for table in tables:
            cursor.execute(f"PRAGMA table_info([{table}])")
            cols = [{"name": r[1], "type": r[2]} for r in cursor.fetchall()]
            cursor.execute(f"SELECT COUNT(*) FROM [{table}]")
            count = cursor.fetchone()[0]
            info["tables"][table] = {"columns": cols, "row_count": count}
        conn.close()
        return info
    except Exception as exc:
        return {"schema_error": f"sqlite error: {exc}"}


def _excel_schema(filepath: Path) -> dict[str, Any]:
    """Extract schema from Excel file."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(filepath, read_only=True)
        info: dict[str, Any] = {"sheets": {}}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(max_row=6, values_only=True))
            if rows:
                info["sheets"][sheet_name] = {
                    "columns": [str(c) for c in rows[0]] if rows else [],
                    "row_count": ws.max_row or 0,
                    "preview": [list(r) for r in rows[1:4]] if len(rows) > 1 else [],
                }
        wb.close()
        return info
    except ImportError:
        return {"schema_error": "openpyxl not available for Excel files"}


def build_data_summary(workdir: Path, max_preview_rows: int = 3) -> str:
    """Build a compact text summary of all discovered data files."""
    files = discover_files(workdir)
    parts: list[str] = []

    if files.get("sample_submission"):
        parts.append("=== SAMPLE SUBMISSION ===")
        schema = infer_schema(files["sample_submission"], max_preview_rows)
        parts.append(_format_schema(schema))

    for fp in files.get("data_files", []):
        if fp == files.get("sample_submission"):
            continue
        parts.append(f"=== {fp.name} ===")
        schema = infer_schema(fp, max_preview_rows)
        parts.append(_format_schema(schema))

    for fp in files.get("info_files", []):
        parts.append(f"=== {fp.name} (info) ===")
        try:
            text = fp.read_text(encoding="utf-8", errors="replace")[:3000]
            parts.append(text)
        except Exception:
            parts.append("(could not read)")

    return "\n\n".join(parts) if parts else "No data files found in workdir."


def _format_schema(schema: dict[str, Any]) -> str:
    """Format a schema dict into a readable string."""
    lines: list[str] = []
    if "columns" in schema:
        lines.append(f"Columns ({len(schema['columns'])}): {schema['columns']}")
    if "row_count" in schema:
        lines.append(f"Rows: {schema['row_count']}")
    if "preview" in schema and schema["preview"]:
        preview = schema["preview"]
        if isinstance(preview, list) and preview and isinstance(preview[0], list):
            for row in preview[:3]:
                lines.append(f"  {row}")
        elif isinstance(preview, list):
            lines.append(f"  Preview: {json.dumps(preview[:2], default=str)[:500]}")
    if "tables" in schema:
        for tname, tinfo in schema["tables"].items():
            lines.append(f"Table '{tname}': {tinfo.get('row_count', '?')} rows, cols: {[c['name'] for c in tinfo.get('columns', [])]}")
    return "\n".join(lines)


def build_llm_prompt(
    task_prompt: str,
    data_summary: str,
    *,
    step_type: str = "plan",
    previous_actions: str | None = None,
    observation: str | None = None,
    out_dir: str | None = None,
    max_chars: int = 12000,
) -> list[dict[str, str]]:
    """Build an LLM message list for the data analysis agent."""
    system_msg = (
        "You are an expert data analyst. You analyze datasets, build models, "
        "and produce structured outputs. Follow these rules:\n"
        "1. Always read and understand the data before acting.\n"
        "2. Use pandas/scikit-learn for data processing and modeling.\n"
        "3. Produce concrete numeric results, not vague descriptions.\n"
        "4. When a sample_submission.csv exists, your submission.csv must match its schema exactly.\n"
        "5. For classification, output the exact label values (True/False, 0/1, etc.) as in the sample.\n"
        "6. For reports, include computed metrics, tables, and quantitative conclusions.\n"
        "7. Respond with executable Python code only when asked to compute something.\n"
        "8. Save output files (submission.csv, REPORT.md, etc.) to the current working directory.\n"
    )

    # Trim data summary to fit budget
    available = max_chars - len(system_msg) - len(task_prompt) - 2000
    if available < 2000:
        available = 2000
    if len(data_summary) > available:
        data_summary = data_summary[:available] + "\n...[truncated]..."

    user_parts = [f"TASK: {task_prompt}\n\nDATA SUMMARY:\n{data_summary}"]

    if out_dir:
        user_parts.append(f"\nOUTPUT DIRECTORY: {out_dir}")

    if previous_actions:
        user_parts.append(f"\nPREVIOUS ACTIONS:\n{previous_actions[:3000]}")

    if observation:
        user_parts.append(f"\nOBSERVATION:\n{observation[:4000]}")

    if step_type == "plan":
        user_parts.append(
            "\nProvide a step-by-step analysis plan. What data processing, "
            "feature engineering, modeling, and output steps should we take?"
        )
    elif step_type == "code":
        out_dir_hint = f" Save output files to: {out_dir}" if out_dir else " Save outputs to the current directory."
        user_parts.append(
            "\nWrite complete Python code to accomplish the next step. "
            "Use pandas, numpy, scikit-learn. Print key results." + out_dir_hint
        )
    elif step_type == "observe":
        user_parts.append(
            "\nBased on the code output above, summarize what you learned "
            "and decide the next action."
        )
    elif step_type == "finalize":
        out_dir_hint = f" Save all output files to: {out_dir}" if out_dir else ""
        user_parts.append(
            "\nWrite final Python code to produce all required output artifacts. "
            "Make sure submission.csv matches the sample schema exactly. "
            "Write REPORT.md with findings." + out_dir_hint
        )

    messages = [
        {"role": "system", "content": system_msg},
        {"role": "user", "content": "\n".join(user_parts)},
    ]
    return messages


def truncate_output(text: str, max_chars: int = 8000) -> str:
    """Truncate text to fit within a character budget."""
    if len(text) <= max_chars:
        return text
    half = max_chars // 2
    return text[:half] + f"\n\n...[truncated {len(text) - max_chars} chars]...\n\n" + text[-half:]
